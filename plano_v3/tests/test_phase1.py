from contextlib import closing
import io
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from app import create_app
from tracker import db

ROOT = Path(__file__).resolve().parents[1]


class PhaseOne(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "test.db"
        self.assertEqual(
            db.initialize(self.path),
            ["0001_tracker", "0002_timer_recovery", "0003_analytics", "0004_topics", "0005_scheduled_break", "0006_rest_time"],
        )
        self.client = create_app(self.path).test_client()

    def tearDown(self):
        self.temp.cleanup()

    def post(self, path, values, status=200):
        r = self.client.post("/api" + path, json=values)
        self.assertEqual(r.status_code, status, r.get_data(as_text=True))
        return r.get_json()

    def course(self, name="Engenharia"):
        return self.post("/courses", {"name": name})["id"]

    def subject(self, course=None, name="Circuitos"):
        return self.post("/subjects", {"course_id": course or self.course(), "name": name})["id"]

    def preview(self, course, content, filename, status=200, **extra):
        r = self.client.post(
            f"/api/courses/{course}/import/preview",
            data={"file": (io.BytesIO(content), filename), **extra},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, status, r.get_data(as_text=True))
        result = r.get_json()
        r.close()
        r.request.close()
        r.request.environ["wsgi.input"].close()
        return result

    def confirm(self, course, items, status=200):
        return self.post(
            f"/courses/{course}/import/confirm", {"confirmed": True, "items": items}, status
        )

    def count(self, table):
        with closing(sqlite3.connect(self.path)) as conn:
            return conn.execute("SELECT COUNT(*) FROM " + table).fetchone()[0]

    def test_schema_idempotent_explicit_path_and_no_legacy_conversion(self):
        self.assertEqual(db.initialize(self.path), [])
        self.assertEqual(db.diagnostics(self.path)["integrity"], "ok")
        missing = Path(self.temp.name) / "missing.db"
        with self.assertRaises(RuntimeError):
            create_app(missing)
        self.assertFalse(missing.exists())
        legacy = Path(self.temp.name) / "legacy.db"
        with closing(sqlite3.connect(legacy)) as c:
            c.execute("CREATE TABLE formacoes(id INTEGER)")
        before = legacy.read_bytes()
        with self.assertRaises(RuntimeError):
            db.initialize(legacy)
        self.assertEqual(before, legacy.read_bytes())

    def test_modified_migration_rejected(self):
        with db.transaction(self.path) as conn:
            conn.execute("UPDATE v3_migrations SET checksum='changed'")
        with self.assertRaises(RuntimeError):
            create_app(self.path)
        with self.assertRaises(RuntimeError):
            db.initialize(self.path)

    def test_model_preview_zero_selection_import_and_restart(self):
        c = self.course()
        result = self.preview(
            c, (ROOT / "static/downloads/modelo_grade_curricular.xlsx").read_bytes(), "grade.xlsx"
        )
        self.assertEqual(result["summary"]["recognized"], 3)
        self.assertEqual(result["summary"]["selected"], 0)
        self.assertEqual(self.count("subjects"), 0)
        self.confirm(c, result["items"], 400)
        result["items"][0]["include"] = True
        result["items"][1]["include"] = True
        summary = self.confirm(c, result["items"])
        self.assertEqual(summary, {"inserted": 2, "updated": 0, "ignored": 1})
        recreated = create_app(self.path).test_client()
        self.assertEqual(len(recreated.get("/api/subjects").json), 2)
        self.assertEqual(self.count("timers"), 0)
        self.assertEqual(db.diagnostics(self.path)["integrity"], "ok")

    def test_courses_keep_homonyms_separate_and_duplicate_update_explicit(self):
        a, b = self.course("Curso A"), self.course("Curso B")
        row = {"name": "Cálculo I", "code": "001", "workload_minutes": 60}
        self.confirm(a, [row])
        self.confirm(b, [row])
        self.assertEqual(self.count("subjects"), 2)
        repeat = self.confirm(a, [{"name": " calculo   i ", "workload_minutes": 120}])
        self.assertEqual(repeat["ignored"], 1)
        self.assertEqual(self.client.get("/api/subjects").json[0]["workload_minutes"], 60)
        self.confirm(a, [{**row, "workload_minutes": 120, "duplicate_action": "update"}])
        self.assertEqual(self.client.get("/api/subjects").json[0]["workload_minutes"], 120)

    def test_batch_rollback_on_invalid_selected_row(self):
        c = self.course()
        self.confirm(c, [{"name": "Válida"}, {"name": "", "include": True}], 400)
        self.assertEqual(self.count("subjects"), 0)
        self.post(
            f"/courses/{c}/import/confirm", {"confirmed": False, "items": [{"name": "A"}]}, 400
        )
        self.assertEqual(self.count("subjects"), 0)

    def test_csv_tsv_txt_and_paste_preserve_accent_code_hours(self):
        c = self.course()
        for extension, separator in [("csv", ";"), ("tsv", "\t"), ("txt", "\t")]:
            source = (
                separator.join(["Disciplina", "Código", "Período", "Carga (h)"])
                + "\n"
                + separator.join(["Álgebra", "001", "1º", "45"])
            )
            result = self.preview(c, source.encode("utf-8-sig"), "grade." + extension)
            self.assertEqual(result["items"][0]["code"], "001")
            self.assertEqual(result["items"][0]["name"], "Álgebra")
            self.assertEqual(result["items"][0]["workload_minutes"], 2700)
        pasted = self.post(f"/courses/{c}/import/preview", {"text": source})
        self.assertEqual(pasted["items"][0]["workload_minutes"], 2700)
        self.assertEqual(self.count("subjects"), 0)

    def test_multiple_sheets_require_choice(self):
        book = Workbook()
        book.active.title = "Grade A"
        book.create_sheet("Grade B")
        for sheet in book:
            sheet.append(["Disciplina", "Carga (h)"])
            sheet.append(["Circuitos", 60])
        content = io.BytesIO()
        book.save(content)
        c = self.course()
        result = self.preview(c, content.getvalue(), "grades.xlsx")
        self.assertTrue(result["requires_sheet_selection"])
        result = self.preview(c, content.getvalue(), "grades.xlsx", sheet="Grade B")
        self.assertEqual(result["selected_sheet"], "Grade B")
        self.assertEqual(len(result["items"]), 1)

    def test_docx_table_and_textual_pdf(self):
        c = self.course()
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        for cell, text in zip(table.rows[0].cells, ["Disciplina", "Código", "Carga (h)"]):
            cell.text = text
        for cell, text in zip(table.rows[1].cells, ["Eletrônica", "001", "30"]):
            cell.text = text
        content = io.BytesIO()
        doc.save(content)
        result = self.preview(c, content.getvalue(), "grade.docx")
        self.assertEqual(result["items"][0]["workload_minutes"], 1800)
        self.assertEqual(result["items"][0]["code"], "001")
        writer = PdfWriter()
        page = writer.add_blank_page(width=595, height=842)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 50 780 Td (UCFC 1) Tj 0 -20 Td (Circuitos 60) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        content = io.BytesIO()
        writer.write(content)
        result = self.preview(c, content.getvalue(), "grade.pdf")
        self.assertEqual(result["items"][0]["workload_minutes"], 3600)

    def test_invalid_empty_large_and_no_text_pdf_are_validation_errors(self):
        c = self.course()
        for name, content in [
            ("invalid.pdf", b"invalid"),
            ("invalid.docx", b"invalid"),
            ("invalid.xlsx", b"invalid"),
            ("empty.csv", b""),
            ("large.txt", b"x" * (5 * 1024 * 1024 + 1)),
        ]:
            result = self.preview(c, content, name, 400)
            self.assertTrue(result["error"])
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        content = io.BytesIO()
        writer.write(content)
        result = self.preview(c, content.getvalue(), "scan.pdf", 400)
        self.assertIn("texto interpretável", result["error"])
        self.assertEqual(self.count("subjects"), 0)

    def test_imported_subject_starts_directly_and_survives_restart(self):
        c = self.course()
        self.confirm(c, [{"name": "Circuitos", "academic_status": "not_available"}])
        subject = self.client.get("/api/subjects").json[0]
        focus = self.post("/timer", {"subject_id": subject["id"], "operation_key": "start-1"})
        self.assertEqual(focus["status"], "running")
        self.assertEqual(focus["subject_id"], subject["id"])
        restarted = create_app(self.path).test_client()
        self.assertEqual(restarted.get("/api/timer").json["id"], focus["id"])
        self.assertEqual(self.count("timers"), 1)
        repeated = self.post("/timer", {"subject_id": subject["id"], "operation_key": "start-1"})
        self.assertEqual(repeated["id"], focus["id"])
        self.post("/timer", {"subject_id": subject["id"], "operation_key": "start-2"}, 409)

    def test_timer_pause_finish_idempotence_and_archive_preserves_history(self):
        c = self.course()
        s = self.subject(c)
        with patch("tracker.timer.now", return_value="2026-10-10T12:00:00+00:00"):
            focus = self.post("/timer", {"subject_id": s, "operation_key": "focus"})
        with patch("tracker.timer.now", return_value="2026-10-10T12:10:00+00:00"):
            paused = self.post(f"/timer/{focus['id']}/pause", {"version": 1})
            self.assertEqual(paused["effective_seconds"], 600)
        with patch("tracker.timer.now", return_value="2026-10-10T12:15:00+00:00"):
            resumed = self.post(f"/timer/{focus['id']}/resume", {"version": paused["version"]})
            self.assertEqual(resumed["break_seconds"], 300)
            self.post(f"/timer/{focus['id']}/pause", {"version": 1}, 409)
        self.post(f"/course/{c}/archive", {"archived": True})
        with patch("tracker.timer.now", return_value="2026-10-10T12:35:00+00:00"):
            finished = self.post(
                f"/timer/{focus['id']}/finish", {"version": resumed["version"], "note": "Revisão"}
            )
            again = self.post(f"/timer/{focus['id']}/finish", {"version": resumed["version"]})
        self.assertEqual(finished["session_id"], again["session_id"])
        self.assertEqual(finished["effective_seconds"], 1800)
        self.assertEqual(finished["break_seconds"], 300)
        self.assertEqual(self.count("study_sessions"), 1)
        self.assertEqual(self.client.get("/api/sessions/recent").json[0]["note"], "Revisão")
        self.assertEqual(self.client.get("/api/subjects").json, [])
        self.post("/timer", {"subject_id": s, "operation_key": "archived"}, 409)
        self.post(f"/course/{c}/archive", {"archived": False})
        self.assertEqual(len(self.client.get("/api/subjects").json), 1)

    def test_cancel_and_free_study_do_not_require_course(self):
        focus = self.post("/timer", {"title": "Leitura livre", "operation_key": "free"})
        self.assertIsNone(focus["subject_id"])
        self.post(f"/timer/{focus['id']}/cancel", {"version": focus["version"]})
        self.assertEqual(self.count("study_sessions"), 0)
        self.assertIsNone(self.client.get("/api/timer").json)

    def test_validation_and_local_request_boundary(self):
        c = self.course()
        self.post("/subjects", {"course_id": c, "name": "A", "workload_minutes": -1}, 400)
        r = self.client.post(
            "/api/courses", json={"name": "x"}, headers={"Origin": "https://other.example"}
        )
        self.assertEqual(r.status_code, 403)
        r = self.client.get("/", headers={"Host": "other.example"})
        self.assertEqual(r.status_code, 403)
        for route in ["/", "/subjects", "/analytics", "/api/diagnostics"]:
            self.assertEqual(self.client.get(route).status_code, 200)

    def test_edit_archive_restore_and_duplicate_validation(self):
        course = self.course()
        subject = self.subject(course)
        edited = self.client.patch(
            f"/api/subjects/{subject}",
            json={"name": "Circuitos II", "code": "002", "workload_minutes": 1800},
        )
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(edited.json["name"], "Circuitos II")
        self.post(f"/subject/{subject}/archive", {"archived": True})
        self.assertEqual(self.client.get("/api/subjects").json, [])
        self.assertEqual(len(self.client.get("/api/subjects?archived=1").json), 1)
        self.post("/subjects", {"course_id": course, "name": "circuitos ii"}, 409)
        self.post(f"/subject/{subject}/archive", {"archived": False})
        row = self.client.get("/api/subjects").json[0]
        self.assertEqual((row["id"], row["code"], row["workload_minutes"]), (subject, "002", 1800))
        self.assertEqual(db.diagnostics(self.path)["integrity"], "ok")

    def test_missing_table_refuses_startup(self):
        with db.transaction(self.path) as conn:
            conn.execute("DROP TABLE study_sessions")
        with self.assertRaisesRegex(RuntimeError, "incompleto"):
            create_app(self.path)


if __name__ == "__main__":
    unittest.main()
